import io
import fitz
from docx import Document

def extract_resume_text(uploaded_file):
    data = uploaded_file.read()
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        doc = fitz.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text("text") for page in doc)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("Unsupported file type.")
