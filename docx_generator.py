from io import BytesIO
from docx import Document
from docx.shared import Pt

def create_docx(text):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        if line.isupper() and len(line) < 60:
            run = p.add_run(line)
            run.bold = True
        else:
            p.add_run(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
